#!/usr/bin/env python3
"""Extract heterogeneous video-recognition workbooks into canonical JSON.

The script is read-only for source xlsx files.  It is intentionally separate
from the workbook builder so the mapping can be inspected and reused for a
future drop of recognition results.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = ROOT / "Data" / "raw" / "video_recognition"

HEADER_ALIASES = {
    "row": {"#", "№", "row", "rownumber", "rowno", "number"},
    "make_model": {"carmakemodel", "makemodel", "vehicle", "make/model", "make / model"},
    "plate": {"licenseplate", "plate"},
    "timestamp": {"timestamp", "videotimestamp"},
    "video_part": {"videopart", "videopartnumber", "partnumber", "videopart№", "videopartno"},
    "video_date": {"videodate"},
    "video_time": {"videotime"},
}

CYR_TO_LAT = str.maketrans({
    "А": "A", "а": "A", "В": "B", "в": "B", "Е": "E", "е": "E",
    "К": "K", "к": "K", "М": "M", "м": "M", "Н": "H", "н": "H",
    "О": "O", "о": "O", "Р": "P", "р": "P", "С": "C", "с": "C",
    "Т": "T", "т": "T", "Х": "X", "х": "X", "І": "I", "і": "I",
    "Ї": "I", "ї": "I", "Є": "E", "є": "E", "Ґ": "G", "ґ": "G",
    "У": "Y", "у": "Y",
})


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def key(value) -> str:
    return re.sub(r"[^a-zа-яіїє0-9#№/]", "", text(value).lower())


def canonical_header(value) -> str | None:
    source = key(value)
    for target, aliases in HEADER_ALIASES.items():
        if source in aliases:
            return target
    return None


def filename_metadata(name: str) -> tuple[str, str]:
    period = "NIGHT" if re.search(r"night|ніч", name, flags=re.I) else "DAY"
    match = re.search(r"(\d{1,2})[._-](\d{1,2})[._-](\d{4})", name)
    if not match:
        raise ValueError(f"Не удалось извлечь дату из имени: {name}")
    day, month, year = map(int, match.groups())
    return date(year, month, day).isoformat(), period


def format_timestamp(value) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}:{value.second:02d}"
    raw = text(value)
    if not raw:
        return ""
    match = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", raw)
    if not match:
        return raw
    first, second, third = match.groups()
    if third is None:
        return f"00:{int(first):02d}:{int(second):02d}"
    return f"{int(first):02d}:{int(second):02d}:{int(third):02d}"


def format_clock(value) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    raw = text(value)
    match = re.search(r"(\d{1,2}):(\d{2})", raw)
    return f"{int(match.group(1)):02d}:{int(match.group(2)):02d}" if match else ""


def format_date(value) -> str:
    """Return an ISO date when a source column contains a recognizable date."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    month_names = {
        "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
        "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
        "січня": 1, "лютого": 2, "березня": 3, "квітня": 4, "травня": 5, "червня": 6,
        "липня": 7, "серпня": 8, "вересня": 9, "жовтня": 10, "листопада": 11, "грудня": 12,
    }
    word_match = re.search(r"(\d{1,2})\s+([а-яіїєґ]+)\s+(\d{4})", raw.lower())
    if word_match and word_match.group(2) in month_names:
        day, year = int(word_match.group(1)), int(word_match.group(3))
        try:
            return date(year, month_names[word_match.group(2)], day).isoformat()
        except ValueError:
            return raw
    match = re.search(r"(\d{1,2})[._/-](\d{1,2})[._/-](\d{4})", raw)
    if not match:
        return raw
    day, month, year = map(int, match.groups())
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return raw


def normalize_plate(value) -> tuple[str, str]:
    raw = text(value)
    if not raw:
        return "", "MISSING"
    normalized = re.sub(r"[^A-Z0-9]", "", raw.upper().translate(CYR_TO_LAT))
    return normalized, "STANDARD" if re.fullmatch(r"[A-Z]{2}\d{4}[A-Z]{2}", normalized) else "CHECK"


def normalize_make(value) -> str:
    raw = text(value)
    if not raw or raw.lower() in {"not specified", "n/a", "-"}:
        return ""
    return re.sub(r"\s+", " ", raw).strip().upper()


def find_header_and_title(ws) -> tuple[int, dict[str, int], str]:
    title_cells = []
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True), start=1):
        values = list(row)
        title_cells.extend(text(value) for value in values if text(value))
        mapping = {}
        for index, value in enumerate(values):
            target = canonical_header(value)
            if target and target not in mapping:
                mapping[target] = index
        if "plate" in mapping and len(mapping) >= 3:
            return row_number, mapping, " | ".join(title_cells)
    raise ValueError(f"Заголовок не найден на листе {ws.title!r}")


def value_at(row, mapping, field):
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def extract_file(path: Path) -> tuple[list[dict], dict]:
    video_date, period = filename_metadata(path.name)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[workbook.sheetnames[0]]
        header_row, mapping, title = find_header_and_title(ws)
        title_time = format_clock(title)
        rows = []
        skipped = 0
        for excel_row, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            plate_raw = text(value_at(values, mapping, "plate"))
            make_raw = text(value_at(values, mapping, "make_model"))
            row_value = text(value_at(values, mapping, "row"))
            if not plate_raw and not make_raw:
                skipped += 1
                continue
            if not plate_raw:
                skipped += 1
                continue
            plate_normalized, plate_status = normalize_plate(plate_raw)
            video_time = format_clock(value_at(values, mapping, "video_time")) or title_time
            part_raw = text(value_at(values, mapping, "video_part"))
            part_match = re.search(r"\d+", part_raw)
            rows.append({
                "row": row_value or str(excel_row - header_row),
                "make_model": make_raw,
                "make_model_normalized": normalize_make(make_raw),
                "plate": plate_raw,
                "plate_normalized": plate_normalized,
                "plate_status": plate_status,
                "video_date": video_date,
                "video_time": video_time,
                "timestamp": format_timestamp(value_at(values, mapping, "timestamp")),
                "video_part_number": part_match.group(0) if part_match else part_raw,
                "recording_period": period,
                "source_file": path.name,
                "source_sheet": ws.title,
                "source_row": excel_row,
                "source_video_date": format_date(value_at(values, mapping, "video_date")),
                "source_video_time": text(value_at(values, mapping, "video_time")),
            })
        return rows, {
            "source_file": path.name,
            "source_sheet": ws.title,
            "video_date": video_date,
            "recording_period": period,
            "header_row": header_row,
            "columns_found": ", ".join(sorted(mapping)),
            "records": len(rows),
            "skipped_rows": skipped,
            "title_time": title_time,
        }
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Standardize video-recognition xlsx files into JSON.")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    source_dir = args.source_dir.resolve()
    files = sorted(path for path in source_dir.glob("*.xlsx") if not path.name.startswith("~$"))
    all_rows, sources, errors = [], [], []
    for path in files:
        try:
            rows, source = extract_file(path)
            all_rows.extend(rows)
            sources.append(source)
        except Exception as error:
            errors.append({"source_file": path.name, "error": str(error)})
    payload = {"observations": all_rows, "sources": sources, "errors": errors}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Files: {len(files)}; observations: {len(all_rows)}; errors: {len(errors)}")
    for error in errors:
        print(f"ERROR {error['source_file']}: {error['error']}")
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
