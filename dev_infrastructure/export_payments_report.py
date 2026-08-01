"""
Экспорт платежей из payments за указанный месяц в Excel-отчёт —
расширенная версия с несколькими листами:

  1. Платежи        — как раньше (дата, квартира, авто, сумма, период, касса)
  2. Сводка по авто  — количество авто по подъездам (1-6) + разбивка по
                        тарифу (День/Ночь/Не определён)
  3. Ведомость       — Квартира | Номер | Тариф (пусто, для заполнения на
                        бумаге) — только авто с неопределённым тарифом
  4. Неполные данные — авто, у которых не хватает режима парковки и/или
                        номера

Использование:
    python export_payments_report.py 2026-07

Файл записывается в Data/exports/ (через config.py).
"""

import sys
import sqlite3
from pathlib import Path
from datetime import datetime

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from config import paths, USE_TEST_DB

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def get_db_path():
    return paths.OSBB_TEST_DB_FILE if USE_TEST_DB else paths.OSBB_DB_FILE


def fetch_payments(conn, month_prefix: str) -> pd.DataFrame:
    rows = conn.execute("""
        SELECT p.id, substr(p.payment_date, 1, 10) AS payment_date, p.apartment_number,
               COALESCE(v.license_plate_normalized, v.license_plate) AS direct_plate,
               p.amount, p.period_code, p.cashbox_code
        FROM payments p
        LEFT JOIN vehicles v ON v.id = p.vehicle_id
        WHERE p.payment_date LIKE ?
        ORDER BY p.payment_date, p.apartment_number
    """, (f"{month_prefix}%",)).fetchall()

    # Тарифы известных режимов — один раз.
    tariff_by_mode = {}
    for mode, code in (("Day", "PARKING_DAY"), ("Night", "PARKING_NIGHT")):
        tariff_row = conn.execute(
            "SELECT amount FROM service_tariffs WHERE service_code=? AND is_active=1 ORDER BY valid_from DESC LIMIT 1",
            (code,),
        ).fetchone()
        tariff_by_mode[mode] = tariff_row[0] if tariff_row else None

    # Группируем НЕразрешённые строки по (квартира, период) — решаем
    # пазл на уровне всех платежей квартиры сразу, не по одной строке.
    # Если число платежей и их суммы (как мультимножество) точно
    # совпадают с полным набором тарифов всех авто квартиры — вопрос
    # закрыт полностью, никакого "?"/"N авто" не нужно: все машины
    # оплачены, какая строка какой соответствует — уже не важно.
    unresolved_idx = [i for i, r in enumerate(rows) if not r["direct_plate"] and r["apartment_number"]]
    groups: dict[tuple, list[int]] = {}
    for i in unresolved_idx:
        key = (rows[i]["apartment_number"], rows[i]["period_code"])
        groups.setdefault(key, []).append(i)

    resolved_plate: dict[int, str] = {}
    for (apt, period), idxs in groups.items():
        vehicles = conn.execute("""
            SELECT COALESCE(v.license_plate_normalized, v.license_plate) AS plate, v.parking_time
            FROM vehicles v JOIN apartments a ON a.id = v.apartment_id
            WHERE a.apartment_number = ?
        """, (apt,)).fetchall()
        vehicle_tariffs = [
            (v["plate"], tariff_by_mode.get(v["parking_time"]))
            for v in vehicles if v["plate"] and tariff_by_mode.get(v["parking_time"])
        ]

        amounts = [float(rows[i]["amount"]) for i in idxs]
        solved_plates = None
        if len(amounts) == len(vehicle_tariffs) and vehicle_tariffs:
            from collections import Counter
            if Counter(amounts) == Counter(t for _, t in vehicle_tariffs):
                solved_plates = sorted(p for p, _ in vehicle_tariffs)

        if solved_plates:
            for i, plate in zip(idxs, solved_plates):
                resolved_plate[i] = plate
        else:
            # Не решилось группой — откат к прежней, по-строчной логике.
            all_plates = [p for p, _ in vehicle_tariffs] or [v["plate"] for v in vehicles if v["plate"]]
            for i in idxs:
                if len(all_plates) == 1:
                    resolved_plate[i] = all_plates[0] + "?"
                elif len(all_plates) > 1:
                    resolved_plate[i] = f"{len(all_plates)}авто"
                else:
                    resolved_plate[i] = ""

    records = []
    for i, r in enumerate(rows):
        plate = r["direct_plate"] or resolved_plate.get(i, "")
        records.append({
            "payment_date": r["payment_date"], "apartment_number": r["apartment_number"],
            "vehicle_plate": plate, "amount": r["amount"],
            "period_code": r["period_code"], "cashbox_code": r["cashbox_code"],
        })
    return pd.DataFrame.from_records(records)


def fetch_entrance_summary(conn) -> pd.DataFrame:
    return pd.read_sql_query(
        """
        SELECT COALESCE(a.entrance_number, a.entrance, 'не указан') AS "Подъезд",
               COUNT(v.id) AS "Авто"
        FROM vehicles v
        JOIN apartments a ON a.id = v.apartment_id
        GROUP BY "Подъезд"
        ORDER BY "Подъезд"
        """,
        conn,
    )


def fetch_tariff_summary(conn) -> dict:
    row = conn.execute("""
        SELECT
            SUM(CASE WHEN parking_time='Day' THEN 1 ELSE 0 END) AS day_count,
            SUM(CASE WHEN parking_time='Night' THEN 1 ELSE 0 END) AS night_count,
            SUM(CASE WHEN parking_time IS NULL THEN 1 ELSE 0 END) AS unspecified_count,
            COUNT(*) AS total
        FROM vehicles
    """).fetchone()
    return {"day": row[0] or 0, "night": row[1] or 0, "unspecified": row[2] or 0, "total": row[3] or 0}


def fetch_worksheet(conn) -> pd.DataFrame:
    df = pd.read_sql_query(
        """
        SELECT a.apartment_number AS "Квартира", COALESCE(v.license_plate_normalized, v.license_plate) AS "Номер"
        FROM vehicles v
        JOIN apartments a ON a.id = v.apartment_id
        WHERE v.parking_time IS NULL
        ORDER BY a.apartment_number
        """,
        conn,
    )
    df["Тариф"] = ""  # пусто намеренно — заполняется на бумаге
    return df


def fetch_incomplete(conn) -> pd.DataFrame:
    return pd.read_sql_query(
        """
        SELECT a.apartment_number AS "Квартира", COALESCE(v.license_plate_normalized, v.license_plate) AS "Номер",
               CASE WHEN v.parking_time IS NULL THEN 'да' ELSE '' END AS "Режим не задан",
               CASE WHEN (v.license_plate_normalized IS NULL OR v.license_plate_normalized = '')
                    AND (v.license_plate IS NULL OR v.license_plate = '')
                    THEN 'да' ELSE '' END AS "Номер не задан"
        FROM vehicles v
        LEFT JOIN apartments a ON a.id = v.apartment_id
        WHERE v.parking_time IS NULL
           OR (v.license_plate_normalized IS NULL OR v.license_plate_normalized = '')
              AND (v.license_plate IS NULL OR v.license_plate = '')
        ORDER BY a.apartment_number
        """,
        conn,
    )


def _style_sheet(ws, headers, df):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, name="Arial")
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name="Arial")
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            [len(str(headers[col_idx - 1]))]
            + [len(str(v)) for v in df.iloc[:, col_idx - 1].astype(str)]
        ) if not df.empty else len(headers[col_idx - 1])
        ws.column_dimensions[col_letter].width = max_len + 3


def build_report(conn, month_prefix: str, out_path: Path):
    payments_df = fetch_payments(conn, month_prefix)
    entrance_df = fetch_entrance_summary(conn)
    tariff = fetch_tariff_summary(conn)
    worksheet_df = fetch_worksheet(conn)
    incomplete_df = fetch_incomplete(conn)

    total = float(payments_df["amount"].sum()) if not payments_df.empty else 0.0
    cash = float(payments_df.loc[payments_df["cashbox_code"] != "BANK", "amount"].sum()) if not payments_df.empty else 0.0
    bank = float(payments_df.loc[payments_df["cashbox_code"] == "BANK", "amount"].sum()) if not payments_df.empty else 0.0

    headers1 = ["Дата платежа", "Квартира", "Номер авто", "Сумма", "Период", "Касса"]
    payments_display = payments_df.rename(columns={
        "payment_date": "Дата платежа", "apartment_number": "Квартира",
        "vehicle_plate": "Номер авто", "amount": "Сумма",
        "period_code": "Период", "cashbox_code": "Касса",
    })[headers1]

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # --- Лист 1: Платежи ---
        payments_display.to_excel(writer, sheet_name="Платежи", index=False)
        ws1 = writer.sheets["Платежи"]
        _style_sheet(ws1, headers1, payments_display)

        summary_row = ws1.max_row + 2
        for i, (label, value) in enumerate([
            ("Итого за " + month_prefix + ":", total),
            ("  из них наличные:", cash),
            ("  из них банк:", bank),
        ]):
            r = summary_row + i
            ws1.cell(row=r, column=1, value=label).font = Font(bold=(i == 0), name="Arial")
            vc = ws1.cell(row=r, column=2, value=value)
            vc.font = Font(bold=(i == 0), name="Arial")
            vc.number_format = "#,##0.00"

        # --- Лист 2: Сводка по авто ---
        entrance_df.to_excel(writer, sheet_name="Сводка по авто", index=False, startrow=0)
        ws2 = writer.sheets["Сводка по авто"]
        _style_sheet(ws2, ["Подъезд", "Авто"], entrance_df)

        tariff_start = ws2.max_row + 2
        ws2.cell(row=tariff_start, column=1, value="Тариф").font = Font(bold=True, name="Arial")
        for i, (label, value) in enumerate([
            ("День", tariff["day"]),
            ("Ночь", tariff["night"]),
            ("Не определён", tariff["unspecified"]),
            ("Всего авто", tariff["total"]),
        ]):
            r = tariff_start + 1 + i
            ws2.cell(row=r, column=1, value=label).font = Font(name="Arial", bold=(label == "Всего авто"))
            ws2.cell(row=r, column=2, value=value).font = Font(name="Arial", bold=(label == "Всего авто"))

        # --- Лист 3: Ведомость (для заполнения на бумаге) ---
        worksheet_df.to_excel(writer, sheet_name="Ведомость", index=False)
        ws3 = writer.sheets["Ведомость"]
        _style_sheet(ws3, ["Квартира", "Номер", "Тариф"], worksheet_df)

        # --- Лист 4: Неполные данные ---
        incomplete_df.to_excel(writer, sheet_name="Неполные данные", index=False)
        ws4 = writer.sheets["Неполные данные"]
        _style_sheet(ws4, ["Квартира", "Номер", "Режим не задан", "Номер не задан"], incomplete_df)


def main():
    if len(sys.argv) < 2:
        print("Использование: python export_payments_report.py ГГГГ-ММ")
        print("Пример: python export_payments_report.py 2026-07")
        return

    month_prefix = sys.argv[1]

    db_path = get_db_path()
    print(f"📁 БД: {db_path}")
    print(f"Период: {month_prefix}")

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        df = fetch_payments(conn, month_prefix)
        print(f"Найдено платежей: {len(df)}")
        if df.empty:
            print("⚠ За этот период платежей не найдено — лист платежей будет пустым.")

        out_dir = paths.EXPORTS_DIR
        out_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = out_dir / f"payments_{month_prefix}_{stamp}.xlsx"

        build_report(conn, month_prefix, out_path)
        print(f"✅ Готово: {out_path}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()