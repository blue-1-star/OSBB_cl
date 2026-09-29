"""Paper remote-control archive; the typed workbook is authoritative."""

from pathlib import Path
import os
import re
import sys
import tempfile

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from admin_console.utils.db import get_conn
from admin_console.utils.session_actor import session_actor
from audit_logger import audit_log

SOURCE = ROOT / "data/raw/typed/Vidacha_Pultiv_In_Order_review.xlsx"


def excel_is_open():
    return SOURCE.with_name("~$" + SOURCE.name).exists()


def load_rows():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Не найдена ведомость: {SOURCE}")
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        sheet = wb.active
        result = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=4, max_col=4, values_only=True), start=4
        ):
            apartment, name, amount, quantity = cells
            if str(name or "").strip().upper() == "ІТОГО:":
                break
            if amount is None and quantity is None and not apartment and not name:
                continue
            result.append({
                "excel_row": row_number,
                "apartment": "" if apartment is None else str(apartment).strip(),
                "name": "" if name is None else str(name).strip(),
                "amount": float(amount or 0),
                "quantity": int(quantity or 0),
            })
        return result
    finally:
        wb.close()


def apartment_sort_key(row):
    """Numeric apartment order; composite labels use their first number."""
    match = re.match(r"^\s*(\d+)", row["apartment"])
    if match:
        return (0, int(match.group(1)), row["excel_row"])
    return (1, 0, row["excel_row"])


def save_row(row_number, apartment, name, actor, evidence):
    if excel_is_open():
        raise ValueError("Файл открыт в Excel. Сохраните и закройте его перед изменением из консоли.")
    apartment, name, evidence = apartment.strip(), name.strip(), evidence.strip()
    if not evidence:
        raise ValueError("Укажите основание изменения.")
    if apartment and not all(part.strip().isdigit() for part in apartment.split(",")):
        raise ValueError("Номер квартиры: цифры; для двух квартир — через запятую.")
    before_mtime = SOURCE.stat().st_mtime_ns
    wb = load_workbook(SOURCE)
    sheet = wb.active
    old_apartment, old_name = sheet.cell(row_number, 1).value, sheet.cell(row_number, 2).value
    changes = []
    for field, old, new in (("Квартира", old_apartment, apartment or None), ("ФИО", old_name, name)):
        if (str(old).strip() if old is not None else "") != (new or ""):
            changes.append((field, old, new))
    if not changes:
        wb.close()
        return False
    sheet.cell(row_number, 1).value = apartment or None
    sheet.cell(row_number, 2).value = name or None
    if SOURCE.stat().st_mtime_ns != before_mtime or excel_is_open():
        wb.close()
        raise ValueError("Файл изменился или был открыт в Excel. Обновите страницу и повторите правку.")
    descriptor, temp_name = tempfile.mkstemp(prefix=".remote-archive-", suffix=".xlsx", dir=SOURCE.parent)
    os.close(descriptor)
    try:
        wb.save(temp_name)
        wb.close()
        if SOURCE.stat().st_mtime_ns != before_mtime or excel_is_open():
            raise ValueError("Файл изменился во время сохранения. Обновите страницу.")
        os.replace(temp_name, SOURCE)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    conn = get_conn()
    try:
        for field, old, new in changes:
            audit_log(conn=conn, operator_id=actor, user_id=actor,
                      actor_type="local_single_user", action_type="remote_archive_excel_corrected",
                      table_name=SOURCE.name, row_id=row_number, field_name=field,
                      old_value=old, new_value=new, source_context=str(SOURCE),
                      comment=evidence, commit=False)
        conn.commit()
    finally:
        conn.close()
    return True


st.set_page_config(page_title="Архив пультов", page_icon="📜", layout="wide")
st.title("📜 Архив выдачи пультов")
st.caption("Источник — Vidacha_Pultiv_In_Order_review.xlsx в data/raw/typed. Сортировка на экране не меняет порядок файла.")
st.info("Это архивный разбор, а не подтверждение оплаты или выдачи. Записи пока не попадают в «Мои заказы» жителя.")
actor = session_actor()
try:
    rows = load_rows()
except (FileNotFoundError, OSError) as exc:
    st.error(str(exc))
    st.stop()
left, right = st.columns(2)
left.metric("Строк в архиве", len(rows))
right.metric("Без номера квартиры", sum(not row["apartment"] for row in rows))
if excel_is_open():
    st.warning("Ведомость открыта в Excel: просмотр доступен, сохранение из консоли временно заблокировано.")

sort_mode = st.radio("Порядок показа", ["Как в бумажной ведомости", "По номеру квартиры"], horizontal=True)
shown_rows = rows if sort_mode == "Как в бумажной ведомости" else sorted(rows, key=apartment_sort_key)
frame = pd.DataFrame([{
    "№": row["excel_row"], "Кв.": row["apartment"] or "—",
    "ФИО": row["name"] or "—", "Сумма": row["amount"],
    "Шт.": row["quantity"],
} for row in shown_rows])
selection = st.dataframe(frame, hide_index=True, width=700, height=580,
    column_config={
        "№": st.column_config.NumberColumn("№", width=60, format="%d"),
        "Кв.": st.column_config.TextColumn("Кв.", width=80),
        "ФИО": st.column_config.TextColumn("ФИО", width=340),
        "Сумма": st.column_config.NumberColumn("грн", width=95, format="%.0f"),
        "Шт.": st.column_config.NumberColumn("Шт.", width=60, format="%d"),
    },
    on_select="rerun", selection_mode="single-row", key=f"remote_archive_file_selection_{sort_mode}")
selected = selection.selection.rows
if not selected:
    st.caption("Нажмите строку ведомости — ниже откроется её карточка.")
    st.stop()
row = shown_rows[selected[0]]
st.subheader(f"Строка Excel {row['excel_row']} · квартира {row['apartment'] or 'не указана'}")
st.write(f"{row['amount']:.0f} грн · {row['quantity']} пульт(а/ов)")
with st.form(f"remote_archive_file_edit_{row['excel_row']}"):
    apartment = st.text_input("Номер квартиры", value=row["apartment"])
    name = st.text_input("ФИО", value=row["name"])
    evidence = st.text_input("Основание изменения", placeholder="Реестр, сообщение, звонок")
    submitted = st.form_submit_button("Сохранить в Excel", type="primary", disabled=excel_is_open())
if submitted:
    try:
        changed = save_row(row["excel_row"], apartment, name, actor, evidence)
    except (ValueError, OSError) as exc:
        st.error(str(exc))
    else:
        if changed:
            st.success("Строка Excel сохранена; действие записано в аудит.")
            st.rerun()
        else:
            st.info("Данные не изменились.")
